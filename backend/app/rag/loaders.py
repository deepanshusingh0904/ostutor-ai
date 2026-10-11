"""Document loaders: PDF, DOCX, TXT, PPTX and Markdown -> list of pages."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".pptx", ".md", ".markdown"}
VIRTUAL_PAGE_CHARS = 3000


@dataclass
class PageText:
    page: int
    text: str


class UnsupportedFileType(ValueError):
    pass


def virtual_pages(text: str, size: int = VIRTUAL_PAGE_CHARS) -> list[PageText]:
    """Split flowing text into pseudo-pages on paragraph boundaries."""
    if "\f" in text:
        parts = [p for p in text.split("\f") if p.strip()]
        return [PageText(i + 1, p) for i, p in enumerate(parts)]
    paragraphs = [p for p in text.replace("\r\n", "\n").split("\n\n") if p.strip()]
    pages: list[PageText] = []
    buffer: list[str] = []
    length = 0
    for para in paragraphs:
        if length + len(para) > size and buffer:
            pages.append(PageText(len(pages) + 1, "\n\n".join(buffer)))
            buffer, length = [], 0
        buffer.append(para)
        length += len(para) + 2
    if buffer:
        pages.append(PageText(len(pages) + 1, "\n\n".join(buffer)))
    return pages


def _load_pdf(path: Path) -> list[PageText]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - damaged page, skip it
            text = ""
        pages.append(PageText(i + 1, text))
    return pages


def _load_docx(path: Path) -> list[PageText]:
    import docx

    document = docx.Document(str(path))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style is not None and paragraph.style.name.lower().startswith("heading"):
            text = f"# {text}"
        blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return virtual_pages("\n\n".join(blocks))


def _load_pptx(path: Path) -> list[PageText]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    pages = []
    for i, slide in enumerate(presentation.slides):
        lines: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        lines.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(notes)
        pages.append(PageText(i + 1, "\n".join(lines)))
    return pages


def _load_text(path: Path) -> list[PageText]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return virtual_pages(text)


def load_document(path: Path) -> list[PageText]:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(f"Unsupported file type: {ext or 'unknown'}")
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".docx":
        return _load_docx(path)
    if ext == ".pptx":
        return _load_pptx(path)
    return _load_text(path)
