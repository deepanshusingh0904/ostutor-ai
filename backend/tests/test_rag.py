from pathlib import Path

from app.rag.chunking import chunk_pages
from app.rag.cleaning import clean_pages, clean_text, join_wrapped_lines
from app.rag.embeddings import HashEmbedder
from app.rag.loaders import PageText, UnsupportedFileType, load_document, virtual_pages
from app.rag.retriever import retrieve


def test_clean_text_and_reflow():
    assert clean_text("a\x00b   c\n\n\n\nd") == "ab c\n\nd"
    assert join_wrapped_lines("the process is\nrunning now") == "the process is running now"
    assert "scheduling" in join_wrapped_lines("sched-\nuling")


def test_repeated_headers_removed():
    pages = [PageText(i, f"Course Notes\nUnique content number {i} here.") for i in range(1, 6)]
    cleaned = clean_pages(pages, "pdf")
    assert all("Course Notes" not in p.text for p in cleaned)


def test_chunking_keeps_pages():
    pages = [PageText(1, "Alpha beta gamma. " * 120), PageText(2, "Delta epsilon zeta. " * 120)]
    chunks = chunk_pages(pages, chunk_size=300, overlap=30)
    assert {c.page for c in chunks} == {1, 2}
    assert all(len(c.text) <= 330 for c in chunks)
    assert [c.index for c in chunks] == list(range(len(chunks)))


def test_virtual_pages_and_loaders(tmp_path: Path):
    text = "\n\n".join(f"Paragraph {i} " + "word " * 100 for i in range(20))
    assert len(virtual_pages(text)) > 1
    f = tmp_path / "notes.txt"
    f.write_text(text)
    assert load_document(f)
    try:
        load_document(tmp_path / "x.exe")
    except UnsupportedFileType:
        pass
    else:
        raise AssertionError("expected UnsupportedFileType")


def test_docx_and_pptx_loaders(tmp_path: Path):
    import docx
    from pptx import Presentation

    d = docx.Document()
    d.add_heading("Deadlock", 1)
    d.add_paragraph("Deadlock needs four conditions to hold at the same time.")
    d.save(tmp_path / "a.docx")
    assert "Deadlock" in load_document(tmp_path / "a.docx")[0].text

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "Paging"
    slide.placeholders[1].text = "Pages map to frames"
    p.save(tmp_path / "a.pptx")
    pages = load_document(tmp_path / "a.pptx")
    assert pages[0].page == 1 and "Paging" in pages[0].text


def test_hash_embedder_similarity():
    e = HashEmbedder()
    a, b, c = e.embed_documents(["cpu scheduling algorithms", "scheduling algorithms for the cpu", "banana bread recipe"])
    dot = lambda x, y: sum(i * j for i, j in zip(x, y))  # noqa: E731
    assert dot(a, b) > dot(a, c)


def test_retrieval_returns_cited_chunks(indexed_document):
    results = retrieve("What is a deadlock and what are the four conditions?")
    assert results
    top = results[0]
    assert top.document_name == "Operating Systems Basics"
    assert top.page >= 1 and 0 < top.score <= 1
    assert "deadlock" in top.text.lower()
